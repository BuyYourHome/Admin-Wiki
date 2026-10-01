"""Read-only Rosebrooks legacy grid snapshot; emits an explicit migration map."""
import collections
import datetime
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import to_excel

path=Path(sys.argv[1])
w=openpyxl.load_workbook(path)
cached=openpyxl.load_workbook(path,data_only=True)
s=w['Carrying'];d=cached['Carrying']
assert not s.tables and 'Carrying - Old' not in w.sheetnames
categories=[('Duke Electric','Duke Electric',2,3),('Mortgage Payment','Mortgage Payment',5,6),
            ('Private Money','Neal Issac',8,9),('Insurance Payments','Insurance Payments',11,12),
            ('Water','City of Raleigh',14,15),('Natural Gas','Dominion Energy',17,18),
            ('HOA','HOA',20,21),('Property Taxes','Property Taxes',23,24)]
records=[];zeros=[];empty=[];totals={}
for category,vendor,dc,ac in categories:
    total=0
    for r in range(3,49):
        date=d.cell(r,dc).value;amount=d.cell(r,ac).value
        source_date=s.cell(r,dc);source_amount=s.cell(r,ac)
        if isinstance(amount,(int,float)) and not isinstance(amount,bool) and amount==0:
            zeros.append({'cell':source_amount.coordinate,'wasFormula':source_amount.data_type=='f'})
            continue
        valid_date=isinstance(date,(datetime.datetime,datetime.date))
        if not valid_date and amount in (None,'',' '):
            empty.append(source_date.coordinate+':'+source_amount.coordinate)
            continue
        assert valid_date,('Unmapped date',source_date.coordinate,date,amount)
        assert amount is None or isinstance(amount,(int,float)) and not isinstance(amount,bool)
        assert date.year>=1900
        record={'category':category,'vendor':vendor,'date':to_excel(date),'amount':amount,
                'dateCell':source_date.coordinate,'amountCell':source_amount.coordinate,
                'dateFormula':source_date.value if source_date.data_type=='f' else None,
                'amountFormula':source_amount.value if source_amount.data_type=='f' else None}
        records.append(record)
        total+=amount or 0
    assert abs(total-(d.cell(49,ac).value or 0))<.000001,(category,total)
    totals[category]=round(total,2)
assert len(records)==24 and sum(r['amount'] is None for r in records)==6
assert abs(sum(totals.values())-10848.20)<.001
result={'project':'Rosebrooks','sourceHash':hashlib.sha256(path.read_bytes()).hexdigest(),
        'sourceName':'20_Project Management - 115 Rosebrooks Dr.xlsm','snapshotDate':'2026-10-01',
        'records':records,'omittedZeroValues':zeros,'emptyPositions':empty,'totals':totals,
        'mode':cached['Profit']['E1'].value,'profitBefore':cached['Profit']['B42'].value,
        'profitTotal':round(sum(totals.values()),2),'keepDocs':True,'docsFormula':w['Docs']['E39'].value,
        'docsValue':to_excel(cached['Docs']['E39'].value) if isinstance(cached['Docs']['E39'].value,datetime.datetime) else cached['Docs']['E39'].value}
Path(sys.argv[2]).write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps({'records':len(records),'blankAmounts':6,'zeroValuesOmitted':len(zeros),
                  'zeroFormulaResults':sum(r['wasFormula'] for r in zeros),'totals':totals,
                  'total':result['profitTotal'],'docsValue':result['docsValue']},indent=2))
