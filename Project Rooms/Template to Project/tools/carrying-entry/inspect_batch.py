"""Read-only, per-workbook Carrying mapping inventory. Never saves a workbook."""
import sys, json, datetime, collections, re
import openpyxl
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel

def value(v):
    return to_excel(v) if isinstance(v,(datetime.datetime,datetime.date)) else v

for path in sys.argv[1:]:
    w=openpyxl.load_workbook(path); cached=openpyxl.load_workbook(path,data_only=True)
    if 'Carrying' not in w.sheetnames:
        print(json.dumps({'file':path,'missing_sheet':True,'sheets':w.sheetnames}))
        continue
    s=w['Carrying']; d=cached['Carrying']
    if 'tblCarryingExpenses' not in s.tables:
        print(json.dumps({'file':path,'missing_table':True,'tables':list(s.tables),'dimensions':s.calculate_dimension(),'cells':[(c.coordinate,str(c.value)) for row in s for c in row if c.value is not None][:150]}))
        continue
    t=s.tables['tblCarryingExpenses']
    left,top,right,bottom=range_boundaries(t.ref)
    headers=[c.name for c in t.tableColumns]
    rows=[dict(zip(headers,[value(d.cell(r,c).value) for c in range(left,right+1)]),row=r) for r in range(top+1,bottom+1)]
    groups=collections.defaultdict(list)
    for r in rows: groups[r['Category']].append(r)
    gridrow=5 if 'ceDate' in w.defined_names else 1
    first=next(r for r in range(gridrow+1,gridrow+8) if s.cell(r,1).data_type=='f' or isinstance(s.cell(r,1).value,(datetime.datetime,datetime.date)) or hasattr(s.cell(r,1).value,'text'))
    footer=next(r for r in range(first+1,left*10) if re.match(r'^=\+?SUM(?:IFS|PRODUCT)?\(',str(s.cell(r,2).value),re.I))
    overrides=[]
    for col in range(1,34,3):
        cat=d.cell(gridrow,col).value
        if cat=='Mortgage Payment paid after Reinstatement':cat='Mortgage Payment'
        records=sorted([r for r in groups[cat] if r['Include']=='Yes'],key=lambda r:(r['Date'] or 0,r['row']))
        for r in range(first,footer):
            for c in (col,col+1):
                cell=s.cell(r,c)
                if cell.value is not None and cell.data_type!='f' and not hasattr(cell.value,'text'):
                    index=r-first
                    overrides.append({'grid':cell.coordinate,'value':value(cell.value),'category':cat,'field':'Date' if c==col else 'Amount','candidate':records[index] if index<len(records) else None})
    p=w['Profit']; pc=cached['Profit']
    print(json.dumps({'file':path,'table':t.ref,'rows':len(rows),'headers':headers,'gridrow':gridrow,'firstDetailRow':first,'footer':footer,
      'categories':{k:{'count':len(v),'zero':sum(isinstance(r['Amount'],(int,float)) and not isinstance(r['Amount'],bool) and r['Amount']==0 and s.cell(r['row'],left+headers.index('Amount')).data_type!='f' for r in v),'blank':sum(r['Amount'] is None for r in v),'sum':sum(r['Amount'] for r in v if isinstance(r['Amount'],(int,float)) and r['Include']=='Yes')} for k,v in groups.items()},
      'blankVendors':sum(not r['Vendor'] or str(r['Vendor']).isspace() for r in rows),
      'overrides':overrides,'profit':{a:{'formula':str(p[a].value),'value':pc[a].value} for a in ['A41','A42','B28','B42','B43','E1','C9','H73','H74','J74','L82','L83']},
      'totals':{d.cell(gridrow,col).coordinate:{'label':d.cell(gridrow,col).value,'total':d.cell(footer,col+1).value} for col in range(1,34,3)},
      'docs':str(w['Docs']['E39'].value),'review':[(str(v.sqref),v.formula1) for v in w['Review'].data_validations.dataValidation],
      'error_counts':dict(collections.Counter(ws.title for ws in cached for row in ws for c in row if c.data_type=='e'))},default=str))
