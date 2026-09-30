"""Read-only XML audit, including hidden content in merged cells. Never resave XLSM."""
import argparse
import json
import posixpath
import re
import zipfile
from collections import Counter
from xml.etree import ElementTree as ET
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'

def normalized(el, ignore=()):
    if el is None: return None
    return (el.tag,tuple(sorted((k,v) for k,v in el.attrib.items() if k not in ignore)),el.text or '',tuple(normalized(c,ignore) for c in el))

class Package:
    def __init__(self,path):
        self.zip=zipfile.ZipFile(path)
        self.wb=ET.fromstring(self.zip.read('xl/workbook.xml'))
        rels={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) for r in ET.fromstring(self.zip.read('xl/_rels/workbook.xml.rels'))}
        self.sheets={s.get('name'):ET.fromstring(self.zip.read(rels[s.get(RID)])) for s in self.wb.find('s:sheets',NS)}
        self.strings=[''.join(e.itertext()) for e in ET.fromstring(self.zip.read('xl/sharedStrings.xml'))]
        self.styles=ET.fromstring(self.zip.read('xl/styles.xml'))
        self.tables={}
        for p in self.zip.namelist():
            if p.startswith('xl/tables/table') and p.endswith('.xml'):
                t=ET.fromstring(self.zip.read(p))
                for el in t.iter():
                    for key in list(el.attrib):
                        if key.endswith('DxfId'):
                            el.set(key,repr(normalized(self.styles.find('s:dxfs',NS)[int(el.get(key))])))
                self.tables[t.get('name')]=normalized(t,('id',))
    def cells(self,sheet):
        result={}
        for c in self.sheets[sheet].findall('s:sheetData/s:row/s:c',NS):
            f=c.find('s:f',NS);v=c.find('s:v',NS)
            val=v.text if v is not None else None
            if c.get('t')=='s' and val is not None: val=self.strings[int(val)]
            if c.get('t')=='inlineStr': val=''.join(c.find('s:is',NS).itertext())
            result[c.get('r')]=(f.text if f is not None else None,val,c.get('s','0'),c.get('t'))
        return result
    def style(self,idx):
        xf=self.styles.find('s:cellXfs',NS)[int(idx)]
        attrs=dict(xf.attrib)
        for key,group in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
            attrs[key]=normalized(self.styles.find('s:'+group,NS)[int(attrs.get(key,'0'))])
        return (attrs,tuple(normalized(c) for c in xf))

def shifted(cell):
    r,c=coordinate_to_tuple(cell.replace('$',''))
    return get_column_letter(c)+str(r+4) if c<=36 and r<=29 else cell

def cut_formula(formula,local=False):
    if not local and 'Carrying!' not in formula and "'Carrying'!" not in formula:
        return formula
    parts=[]
    for token in Tokenizer('='+formula).items:
        v=token.value
        if token.subtype=='RANGE':
            prefix=''
            if '!' in v: prefix,v=v.rsplit('!',1);prefix+='!'
            eligible=(prefix in ("Carrying!","'Carrying'!")) or (not prefix and local)
            if eligible and re.fullmatch(r'\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?',v):
                def sub(m):
                    col,dollar,row=m.groups()
                    target=shifted(col+row)
                    nr=re.search(r'\d+$',target).group()
                    return col+dollar+nr
                v=re.sub(r'(\$?[A-Z]+)(\$?)(\d+)',sub,v)
            v=prefix+v
        parts.append(v)
    return ''.join(parts)

p=argparse.ArgumentParser();p.add_argument('before');p.add_argument('after');args=p.parse_args()
a,b=Package(args.before),Package(args.after)
issues=[];style_issues=[];dependent=[];totals={};errors_before=set();errors_after=set()
for sheet in a.sheets:
    old,new=a.cells(sheet),b.cells(sheet)
    for addr,cell in old.items():
        row,col=coordinate_to_tuple(addr)
        local=sheet=='Carrying' and col<=36 and row<=25
        target=shifted(addr) if local else addr
        after=new.get(target,(None,None,'0',None))
        expected=cut_formula(cell[0],local) if cell[0] else cell[1]
        actual=after[0] if after[0] else after[1]
        if not (sheet=='Carrying' and col<=36 and 26<=row<=29):
            if expected!=actual:issues.append([sheet,addr,target,expected,actual])
            if a.style(cell[2])!=b.style(after[2]):style_issues.append([sheet,addr,target])
        if cell[0] and cell[0]!=after[0] and sheet!='Carrying':dependent.append([sheet,addr,cell[0],after[0]])
        if cell[3]=='e':errors_before.add((sheet,target,cell[1]))
        if local and row==25 and cell[0]:
            totals[addr+' -> '+target]=[cell[1],after[1]]
            if cell[1]!=after[1]:issues.append(['total',addr,cell[1],after[1]])
    for addr,cell in new.items():
        if cell[3]=='e':errors_after.add((sheet,addr,cell[1]))
        if sheet!='Carrying' and addr not in old and (cell[0] is not None or cell[1] is not None):issues.append(['new cell',sheet,addr])
    for tag in ['pageMargins','pageSetup','printOptions','cols']:
        if normalized(a.sheets[sheet].find('s:'+tag,NS))!=normalized(b.sheets[sheet].find('s:'+tag,NS)):issues.append([tag,sheet])
    merges={m.get('ref') for m in a.sheets[sheet].findall('s:mergeCells/s:mergeCell',NS)}
    target_merges={m.get('ref') for m in b.sheets[sheet].findall('s:mergeCells/s:mergeCell',NS)}
    for merge in merges:
        expected=':'.join(shifted(c) for c in merge.split(':')) if sheet=='Carrying' else merge
        if expected not in target_merges:issues.append(['missing merge',sheet,merge,expected])
names_a={n.get('name'):(n.text or '') for n in a.wb.findall('s:definedNames/s:definedName',NS)}
names_b={n.get('name'):(n.text or '') for n in b.wb.findall('s:definedNames/s:definedName',NS)}
for name,formula in names_a.items():
    if cut_formula(formula)!=names_b.get(name):issues.append(['name',name,formula,names_b.get(name)])
table_changes=[name for name,t in a.tables.items() if b.tables.get(name)!=t]
external_changes=[name for name in a.zip.namelist() if name.startswith('xl/externalLinks/') and (name not in b.zip.namelist() or a.zip.read(name)!=b.zip.read(name))]
ctrl_a=Counter(normalized(ET.fromstring(a.zip.read(n))) for n in a.zip.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
ctrl_b=Counter(normalized(ET.fromstring(b.zip.read(n))) for n in b.zip.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
missing_controls=list((ctrl_a-ctrl_b).elements())
result={'issue_count':len(issues),'issues':issues[:25],'style_issue_count':len(style_issues),'style_issues':style_issues[:30],'dependent_formula_changes':dependent,'totals':totals,'new_errors':sorted(errors_after-errors_before),'existing_errors':len(errors_before),'table_changes':table_changes,'external_link_changes':external_changes,'missing_control_count':len(missing_controls),'missing_controls':missing_controls[:2]}
print(json.dumps(result,indent=2))
if issues or style_issues or errors_after-errors_before or table_changes or external_changes or missing_controls:raise SystemExit(1)
