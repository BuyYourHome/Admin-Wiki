"""Read-only preservation audit for native Carrying zero-row deletion."""
import sys, json, zipfile, posixpath
from xml.etree import ElementTree as ET
import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter, coordinate_to_tuple

N={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
def load(path):
    w=openpyxl.load_workbook(path,keep_vba=True);z=zipfile.ZipFile(path)
    assert z.testzip() is None
    rels={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    wb=ET.fromstring(z.read('xl/workbook.xml'));ss=[''.join(x.itertext()) for x in ET.fromstring(z.read('xl/sharedStrings.xml'))]
    sheets={s.get('name'):ET.fromstring(z.read(rels[s.get(R)])) for s in wb.find('s:sheets',N)}
    cells={}
    for name,s in sheets.items():
        data={}
        for c in s.findall('s:sheetData/s:row/s:c',N):
            f=c.find('s:f',N);v=c.find('s:v',N);value=v.text if v is not None else None
            if c.get('t')=='s' and value is not None:value=ss[int(value)]
            if c.get('t')=='inlineStr':value=''.join(c.find('s:is',N).itertext())
            formula=f.text if f is not None else None
            if f is not None and formula is None and f.get('t')=='shared':
                formula=w[name][c.get('r')].value.removeprefix('=')
            data[c.get('r')]=(formula,value,c.get('t'),int(c.get('s','0')))
        cells[name]=data
    return w,z,cells

def style(w,i):
    st=w._cell_styles[i]
    fmt=w._number_formats[st.numFmtId-164] if st.numFmtId>=164 else st.numFmtId
    return (w._fonts[st.fontId],w._fills[st.fillId],w._borders[st.borderId],w._alignments[st.alignmentId],w._protections[st.protectionId],fmt)

a,az,ac=load(sys.argv[1]);b,bz,bc=load(sys.argv[2]);issues=[]
table=a['Carrying'].tables['tblCarryingExpenses'];left,top,right,bottom=range_boundaries(table.ref)
amount=left+[c.name for c in table.tableColumns].index('Amount')
removed=[];mapping={};dest=top
for r in range(top+1,bottom+1):
    c=a['Carrying'].cell(r,amount);v=c.value
    if c.data_type!='f' and isinstance(v,(int,float)) and not isinstance(v,bool) and v==0:removed.append(r)
    else:dest+=1;mapping[r]=dest
assert list(a.sheetnames)==list(b.sheetnames)
for name,cells in ac.items():
    new=bc[name]
    for addr,value in cells.items():
        r,c=coordinate_to_tuple(addr);target=addr
        in_body=name=='Carrying' and left<=c<=right and top<r<=bottom
        if in_body:
            if r in removed:continue
            target=get_column_letter(c)+str(mapping[r])
        actual=new.get(target,(None,None,None,0))
        expected_value=value[0] if value[0] is not None else value[1]
        actual_value=actual[0] if actual[0] is not None else actual[1]
        if expected_value!=actual_value:issues.append([name,addr,target,'content'])
        if style(a,value[3])!=style(b,actual[3]):issues.append([name,addr,target,'style'])
    for addr,value in new.items():
        r,c=coordinate_to_tuple(addr)
        if name=='Carrying' and left<=c<=right and dest<r<=bottom:
            if value[0] is not None or value[1] is not None:issues.append(['table tail not cleared',addr])
        elif addr not in cells and (value[0] is not None or value[1] is not None):issues.append(['unexpected cell',name,addr])
    for attr in ('merged_cells','page_setup','page_margins','print_options','data_validations'):
        if getattr(a[name],attr)!=getattr(b[name],attr):issues.append(['layout',name,attr])
    for tn,t in a[name].tables.items():
        expected=a[name].tables[tn].ref
        if tn=='tblCarryingExpenses':expected=f'{get_column_letter(left)}{top}:{get_column_letter(right)}{dest}'
        if b[name].tables[tn].ref!=expected:issues.append(['table',tn])
for name,n in a.defined_names.items():
    if name not in b.defined_names or n.attr_text!=b.defined_names[name].attr_text:issues.append(['name',name])
for addr in ['B29','E29','H29','K29','N29','Q29','T29','W29','Z29','AC29','AF29','AI29']:
    if ac['Carrying'][addr][1]!=bc['Carrying'][addr][1]:issues.append(['total',addr])
if ac['Profit']['B43'][1]!=bc['Profit']['B43'][1]:issues.append(['Profit total'])
old_errors={(s,k,v[1]) for s,rows in ac.items() for k,v in rows.items() if v[2]=='e'}
new_errors={(s,k,v[1]) for s,rows in bc.items() for k,v in rows.items() if v[2]=='e'}
if new_errors-old_errors:issues.append(['new errors',sorted(new_errors-old_errors)])
vba_changed=az.read('xl/vbaProject.bin')!=bz.read('xl/vbaProject.bin')
if vba_changed and '--vba-source-verified' not in sys.argv:issues.append(['VBA binary changed: inspect native source'])
old_controls=sorted(az.read(n) for n in az.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
new_controls=sorted(bz.read(n) for n in bz.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
if old_controls!=new_controls:issues.append(['controls changed'])
if any(n.startswith('xl/externalLinks/') for n in bz.namelist()):issues.append(['external links'])
print(json.dumps({'removed':len(removed),'remaining':dest-top,'vba_binary_changed':vba_changed,'issues':issues},indent=2))
if issues:sys.exit(1)
