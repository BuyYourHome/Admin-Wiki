"""Independent read-only audit of the approved old-grid value snapshot replacement."""
import collections
import json
import posixpath
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import coordinate_to_tuple,get_column_letter

NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RID='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
class Book:
    def __init__(self,path):
        self.w=openpyxl.load_workbook(path,keep_vba=True);self.z=zipfile.ZipFile(path)
        assert self.z.testzip() is None
        rels={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) for r in ET.fromstring(self.z.read('xl/_rels/workbook.xml.rels'))}
        wb=ET.fromstring(self.z.read('xl/workbook.xml'))
        strings=[''.join(x.itertext()) for x in ET.fromstring(self.z.read('xl/sharedStrings.xml'))]
        self.cells={}
        for sheet in wb.find('s:sheets',NS):
            xml=ET.fromstring(self.z.read(rels[sheet.get(RID)]));cells={};anchors={}
            for c in xml.findall('s:sheetData/s:row/s:c',NS):
                f=c.find('s:f',NS);v=c.find('s:v',NS);value=v.text if v is not None else None
                if c.get('t')=='s' and value is not None:value=strings[int(value)]
                if c.get('t')=='inlineStr':value=''.join(c.find('s:is',NS).itertext())
                cells[c.get('r')]=(f.text if f is not None else None,value,c.get('t'),int(c.get('s','0')))
                if f is not None and f.get('t')=='shared' and f.text:anchors[f.get('si')]=(c.get('r'),f.text)
            for c in xml.findall('s:sheetData/s:row/s:c',NS):
                f=c.find('s:f',NS)
                if f is not None and f.get('t')=='shared' and not f.text:
                    origin,text=anchors[f.get('si')]
                    cells[c.get('r')]=(Translator('='+text,origin=origin).translate_formula(c.get('r'))[1:],*cells[c.get('r')][1:])
            self.cells[sheet.get('name')]=cells
    def style(self,i):
        s=self.w._cell_styles[i];fmt=self.w._number_formats[s.numFmtId-164] if s.numFmtId>=164 else s.numFmtId
        return (self.w._fonts[s.fontId],self.w._fills[s.fillId],self.w._borders[s.borderId],self.w._alignments[s.alignmentId],self.w._protections[s.protectionId],fmt)

def mapped(sheet,addr):
    r,c=coordinate_to_tuple(addr.replace('$',''))
    return get_column_letter(c)+str(r+1 if sheet=='Profit' and r>=41 else r)

def formula(text,sheet):
    out=[]
    for t in Tokenizer('='+text).items:
        value=t.value
        if t.subtype=='RANGE':
            prefix,ref=value.rsplit('!',1) if '!' in value else ('',value)
            target=prefix.strip("'") if prefix else sheet
            if re.fullmatch(r'\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?',ref):
                def sub(m):
                    dc,c,dr,r=m.groups();nr,nc=coordinate_to_tuple(mapped(target,c+r))
                    return dc+get_column_letter(nc)+dr+str(nr)
                ref=re.sub(r'(\$?)([A-Z]+)(\$?)(\d+)',sub,ref)
            if prefix.strip("'")=='Carrying':prefix="'Carrying - Old'"
            value=(prefix+'!' if prefix else '')+ref
        out.append(value)
    return ''.join(out)

def canonical(text):
    if text is None:return None
    text=''.join(p if p.startswith('"') else re.sub(r'\s+',' ',p) for p in re.findall(r'"(?:[^"]|"")*"|[^"]+',text))
    return text.replace('_xlfn.DAYS(','DAYS(')

a,b,p=map(Book,sys.argv[1:4]);mapping=json.load(open(sys.argv[4],encoding='utf-8-sig'))
issues=[];old_errors=set();new_errors=set()
expected_sheets=[]
for name in a.w.sheetnames:expected_sheets.extend(['Carrying - Old','Carrying'] if name=='Carrying' else [name])
assert b.w.sheetnames==expected_sheets
total_cols=['B','E','H','K','N','Q','T','W','Z','AC','AI','AF']
profit_edits={'B'+str(31+i):'Carrying!'+col+'29/Profit!$B$28' for i,col in enumerate(total_cols)}
profit_edits.update({'B43':'B28*SUM(B31:B42)','L83':'IF(H83=0,0,+J83/H83*$F$75/DAYS(J75,H75))'})
for name,old in a.cells.items():
    newname='Carrying - Old' if name=='Carrying' else name;new=b.cells[newname];positions=set()
    for addr,before in old.items():
        dst=mapped(name,addr);positions.add(dst);after=new.get(dst,(None,None,None,0))
        expected=formula(before[0],name) if before[0] is not None else before[1]
        if name=='Profit' and dst in profit_edits:expected=profit_edits[dst]
        actual=after[0] if after[0] is not None else after[1]
        if canonical(expected)!=canonical(actual):issues.append(['content',name,addr,dst,expected,actual])
        if a.style(before[3])!=b.style(after[3]):issues.append(['style',name,addr,dst])
        if before[2]=='e':old_errors.add((newname,dst,before[1]))
        if name=='Carrying' and before[1]!=after[1]:issues.append(['old grid result changed',addr,before[1],after[1]])
    for dst,value in new.items():
        if dst not in positions and (value[0] is not None or value[1] is not None) and not(name=='Profit' and coordinate_to_tuple(dst)[0]==41):issues.append(['unexpected cell',name,dst])
    for merge in a.w[name].merged_cells.ranges:
        expected=':'.join(mapped(name,x) for x in str(merge).split(':'))
        if expected not in b.w[newname].merged_cells:issues.append(['merge',name,str(merge)])
    for attr in ('page_setup','page_margins','print_options','freeze_panes'):
        if getattr(a.w[name],attr)!=getattr(b.w[newname],attr):issues.append([attr,name])
    for tn in a.w[name].tables:
        src=a.w[name].tables[tn];dst=b.w[newname].tables.get(tn)
        if dst is None or dst.ref!=':'.join(mapped(name,x) for x in src.ref.split(':')) or [c.name for c in src.tableColumns]!=[c.name for c in dst.tableColumns]:issues.append(['table',name,tn])
    for col,d in a.w[name].column_dimensions.items():
        other=b.w[newname].column_dimensions.get(col)
        if other is None or (d.width,d.hidden,d.outlineLevel)!=(other.width,other.hidden,other.outlineLevel):issues.append(['column dimensions',name,col])
    for row,d in a.w[name].row_dimensions.items():
        nr=row+1 if name=='Profit' and row>=41 else row;other=b.w[newname].row_dimensions[nr]
        if (d.height,d.hidden,d.outlineLevel)!=(other.height,other.hidden,other.outlineLevel):issues.append(['row dimensions',name,row])
for name,n in a.w.defined_names.items():
    after=b.w.defined_names.get(name)
    if after is None or canonical(formula(n.attr_text,''))!=canonical(after.attr_text):issues.append(['name',name])
for name,cells in b.cells.items():
    for addr,v in cells.items():
        if v[2]=='e':new_errors.add((name,addr,v[1]))
if new_errors-old_errors:issues.append(['new errors',sorted(new_errors-old_errors)])
table=b.w['Carrying'].tables['tblCarryingExpenses']
assert table.ref=='AL2:AV26'
expected_headers=['Include','Category','Date','Vendor','Description','Amount','Source','Invoice #','Source File','Status','Notes']
assert [c.name for c in table.tableColumns]==expected_headers
for index,record in enumerate(mapping['records'],3):
    values=[b.cells['Carrying'].get(get_column_letter(c)+str(index),(None,None))[1] for c in range(38,49)]
    expected=['Yes',record['category'],format(record['date'],'.15g'),record['vendor'],'Legacy '+record['category']+' entry',None if record['amount'] is None else format(record['amount'],'.15g'),'Legacy Grid Snapshot',None,mapping['sourceName'],'Missing Data' if record['amount'] is None else 'Migrated Snapshot']
    for i,v in enumerate(expected):
        if i in (2,5) and v is not None and values[i] is not None:
            if abs(float(v)-float(values[i]))>0.0000001:issues.append(['record numeric value',index,i,v,values[i]])
            continue
        if values[i] not in (v,'' if v is None else v):issues.append(['record',index,i,v,values[i]])
    if 'Carrying - Old!'+record['dateCell']+':'+record['amountCell'] not in values[10]:issues.append(['provenance',index])
    if any(b.cells['Carrying'].get(get_column_letter(c)+str(index),(None,))[0] for c in range(38,49)):issues.append(['record formula',index])
for addr in ('A2','C2','G2','L2','R2','U2','A4','D4','I4','O4','U4','Z3'):
    if p.style(p.cells['Carrying'][addr][3])!=b.style(b.cells['Carrying'][addr][3]):issues.append(['form style',addr])
for attr in ('page_setup','page_margins','print_options','freeze_panes','print_area'):
    if getattr(p.w['Carrying'],attr)!=getattr(b.w['Carrying'],attr):issues.append(['prototype layout',attr])
if set(map(str,p.w['Carrying'].merged_cells.ranges))!=set(map(str,b.w['Carrying'].merged_cells.ranges)):
    issues.append(['prototype merges'])
for col,d in p.w['Carrying'].column_dimensions.items():
    other=b.w['Carrying'].column_dimensions.get(col)
    if other is None or (d.width,d.hidden,d.outlineLevel)!=(other.width,other.hidden,other.outlineLevel):issues.append(['prototype column',col])
for row in range(1,30):
    if p.w['Carrying'].row_dimensions[row].height!=b.w['Carrying'].row_dimensions[row].height:issues.append(['prototype row height',row])
for addr,v in p.cells['Carrying'].items():
    row,col=coordinate_to_tuple(addr)
    if row<=29 and col<=36 and p.style(v[3])!=b.style(b.cells['Carrying'].get(addr,(None,None,None,0))[3]):issues.append(['prototype grid style',addr])
for addr,v in b.cells['Carrying'].items():
    row,col=coordinate_to_tuple(addr)
    if row>26 and 38<=col<=48 and (v[0] is not None or v[1] is not None):issues.append(['template record residue',addr])
for addr in ('A2','C2','G2','L2','R2','A4','I4','O4','Z3'):
    if b.cells['Carrying'].get(addr,(None,None))[1] is not None:issues.append(['template input residue',addr])
for col in (1,2,4,5,7,8,10,11,13,14,16,17,19,20,22,23,25,26,28,29,31,32,34,35):
    for row in range(8,29):
        addr=get_column_letter(col)+str(row)
        expected=p.cells['Carrying'][addr][0].replace('VALUE(tblCarryingExpenses[Amount])','tblCarryingExpenses[Amount]')
        if canonical(expected)!=canonical(b.cells['Carrying'][addr][0]):issues.append(['grid formula',addr])
for row in range(8,29):
    if b.cells['Carrying'].get('AA'+str(row),(None,None))[:2]!=(None,None):issues.append(['template escrow residue',row])
for scope in [b.w,*b.w.worksheets]:
    for n in scope.defined_names.values():
        if re.search(r'\[.*\.xls|\.xls[mxb]?\x27?!',n.attr_text,re.I):issues.append(['external name',n.name])
if any(n.startswith('xl/externalLinks/') for n in b.z.namelist()):issues.append(['external link package residue'])
def controls(book):return collections.Counter(ET.tostring(ET.fromstring(book.z.read(n))) for n in book.z.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
if controls(a)-controls(b):issues.append(['original controls missing'])
print(json.dumps({'issueCount':len(issues),'issues':issues[:20],'oldErrors':len(old_errors),'remainingErrors':len(new_errors),'removedErrors':sorted(old_errors-new_errors),'records':24,'oldSheetPreserved':not any(x[0]=='old grid result changed' for x in issues),'profitTotal':b.cells['Profit']['B43'][1]},indent=2))
sys.exit(bool(issues))
