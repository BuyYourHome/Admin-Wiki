"""Read-only before/after audit for the Tensity Vendor prefill pilot."""
import json
import posixpath
import sys
import zipfile
from xml.etree import ElementTree as ET
import openpyxl

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def load(path):
    book = openpyxl.load_workbook(path, keep_vba=True)
    archive = zipfile.ZipFile(path)
    assert archive.testzip() is None
    rels = {r.get('Id'): posixpath.normpath('xl/' + r.get('Target'))
            for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
    strings = [''.join(x.itertext()) for x in ET.fromstring(archive.read('xl/sharedStrings.xml'))]
    wb = ET.fromstring(archive.read('xl/workbook.xml'))
    cells = {}
    for sheet in wb.find('s:sheets', NS):
        name = sheet.get('name')
        cells[name] = {}
        xml = ET.fromstring(archive.read(rels[sheet.get(RID)]))
        for c in xml.findall('s:sheetData/s:row/s:c', NS):
            f, v = c.find('s:f', NS), c.find('s:v', NS)
            value = v.text if v is not None else None
            if c.get('t') == 's' and value is not None:
                value = strings[int(value)]
            if c.get('t') == 'inlineStr':
                value = ''.join(c.find('s:is', NS).itertext())
            formula = f.text if f is not None else None
            if f is not None and formula is None and f.get('t') == 'shared':
                formula = book[name][c.get('r')].value.removeprefix('=')
            cells[name][c.get('r')] = (formula, value, c.get('t'), int(c.get('s', '0')))
    return book, archive, cells


def style(book, index):
    s = book._cell_styles[index]
    fmt = book._number_formats[s.numFmtId-164] if s.numFmtId >= 164 else s.numFmtId
    return (book._fonts[s.fontId], book._fills[s.fillId], book._borders[s.borderId],
            book._alignments[s.alignmentId], book._protections[s.protectionId], fmt)


def table_xml(book, table):
    tree = table.to_tree()
    for element in tree.iter():
        for key, value in list(element.attrib.items()):
            if key.lower().endswith('dxfid'):
                element.set(key, ET.tostring(book._differential_styles.styles[int(value)].to_tree()).decode())
    return ET.tostring(tree)


a, az, ac = load(sys.argv[1])
b, bz, bc = load(sys.argv[2])
issues = []
assert a.sheetnames == b.sheetnames
for name, old in ac.items():
    for addr in old.keys() | bc[name].keys():
        before = old.get(addr, (None, None, None, 0))
        after = bc[name].get(addr, (None, None, None, 0))
        helper = name == 'Carrying' and addr in ['AX' + str(i) for i in range(1, 8)]
        vendor = name == 'Carrying' and addr in ['AO' + str(i) for i in range(3, 44)]
        expected = before[0] if before[0] is not None else before[1]
        actual = after[0] if after[0] is not None else after[1]
        if vendor:
            assert expected in (None, '', ' ')
            expected = ac[name]['AM' + addr[2:]][1]
        if not helper and expected != actual:
            issues.append([name, addr, 'content', expected, actual])
        if not helper and style(a, before[3]) != style(b, after[3]):
            issues.append([name, addr, 'style'])
    for attr in ('merged_cells', 'page_setup', 'page_margins', 'print_options', 'freeze_panes'):
        if getattr(a[name], attr) != getattr(b[name], attr):
            issues.append([name, attr])
    for tn in a[name].tables:
        if table_xml(a, a[name].tables[tn]) != table_xml(b, b[name].tables[tn]):
            issues.append([name, tn, 'table definition'])
    old_dv = [ET.tostring(d.to_tree()) for d in a[name].data_validations.dataValidation]
    new_dv = [ET.tostring(d.to_tree()) for d in b[name].data_validations.dataValidation]
    for dv in old_dv:
        if dv not in new_dv: issues.append([name, 'validation removed'])
    if len(new_dv)-len(old_dv) != (1 if name == 'Carrying' else 0):
        issues.append([name, 'unexpected validations'])
for key, value in a.defined_names.items():
    expected = 'Carrying!$Z$3' if key == 'ceFeedback' else value.attr_text
    if key not in b.defined_names or b.defined_names[key].attr_text != expected:
        issues.append(['name', key])
# Excel records the LET variable as a hidden compatibility name.
assert set(b.defined_names)-set(a.defined_names) == {'ceVendorList', '_xlpm.v'}
assert b.defined_names['_xlpm.v'].hidden and b.defined_names['_xlpm.v'].attr_text == '#NAME?'
old_errors = {(s, c, v[1]) for s, rows in ac.items() for c, v in rows.items() if v[2] == 'e'}
new_errors = {(s, c, v[1]) for s, rows in bc.items() for c, v in rows.items() if v[2] == 'e'}
if new_errors != old_errors: issues.append(['errors changed'])
assert not any(n.startswith('xl/externalLinks/') for n in bz.namelist())
assert ac['Profit']['B43'][1] == bc['Profit']['B43'][1]
print(json.dumps({'issues': issues, 'vendors_filled': 41, 'existing_errors': len(old_errors),
                  'profit_unchanged': bc['Profit']['B43'][1]}, indent=2))
sys.exit(bool(issues))
