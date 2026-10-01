"""Independent read-only Outrigger mapping audit; never writes a workbook."""
import sys, re, json, zipfile, posixpath
from collections import Counter
from xml.etree import ElementTree as ET
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'

def mapped(sheet, addr):
    r, c = coordinate_to_tuple(addr.replace('$',''))
    if sheet == 'Profit' and r >= 41: r += 1
    if sheet == 'Carrying':
        if c <= 33 and r <= 29: r += 4
        elif 35 <= c <= 45 and r <= 134: c += 3
    return get_column_letter(c)+str(r)

def formula(text, sheet):
    out=[]
    for token in Tokenizer('='+text).items:
        v=token.value
        if token.subtype == 'RANGE':
            prefix, ref = v.rsplit('!',1) if '!' in v else ('',v)
            target=prefix.strip("'") if prefix else sheet
            if re.fullmatch(r'\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?',ref):
                def sub(m):
                    dc,c,dr,r=m.groups(); dst=mapped(target,c+r)
                    nr,nc=coordinate_to_tuple(dst)
                    return dc+get_column_letter(nc)+dr+str(nr)
                ref=re.sub(r'(\$?)([A-Z]+)(\$?)(\d+)',sub,ref)
            v=(prefix+'!' if prefix else '')+ref
        out.append(v)
    return ''.join(out)

def formula_equal(a,b):
    if a is None or b is None: return a==b
    # Excel serializes DAYS with a compatibility prefix; tokenizer collapses
    # insignificant whitespace but preserves strings and intersection tokens.
    def tokens(s):
        return [(t.type,t.subtype,t.value.replace('_xlfn.DAYS(', 'DAYS(')) for t in Tokenizer('='+s).items]
    return tokens(a)==tokens(b)

class Book:
    def __init__(self,path):
        self.w=openpyxl.load_workbook(path,keep_vba=True)
        self.z=zipfile.ZipFile(path); assert self.z.testzip() is None
        wb=ET.fromstring(self.z.read('xl/workbook.xml'))
        rels={r.get('Id'):posixpath.normpath('xl/'+r.get('Target')) for r in ET.fromstring(self.z.read('xl/_rels/workbook.xml.rels'))}
        self.sheets={s.get('name'):ET.fromstring(self.z.read(rels[s.get(RID)])) for s in wb.find('s:sheets',NS)}
        strings=[''.join(e.itertext()) for e in ET.fromstring(self.z.read('xl/sharedStrings.xml'))]
        self.cells={}
        for name,s in self.sheets.items():
            cells={}
            for c in s.findall('s:sheetData/s:row/s:c',NS):
                f=c.find('s:f',NS);v=c.find('s:v',NS); val=v.text if v is not None else None
                if c.get('t')=='s' and val is not None: val=strings[int(val)]
                if c.get('t')=='inlineStr': val=''.join(c.find('s:is',NS).itertext())
                cells[c.get('r')]=(f.text if f is not None else None,val,c.get('t'),int(c.get('s','0')))
            anchors={c.find('s:f',NS).get('si'):c for c in s.findall('s:sheetData/s:row/s:c',NS) if c.find('s:f',NS) is not None and c.find('s:f',NS).get('t')=='shared' and c.find('s:f',NS).text}
            for c in s.findall('s:sheetData/s:row/s:c',NS):
                f=c.find('s:f',NS)
                if f is not None and f.get('t')=='shared' and not f.text:
                    anchor=anchors[f.get('si')]
                    expanded=Translator('='+anchor.find('s:f',NS).text,origin=anchor.get('r')).translate_formula(c.get('r'))[1:]
                    cells[c.get('r')]=(expanded,*cells[c.get('r')][1:])
            self.cells[name]=cells
    def style(self,i):
        st=self.w._cell_styles[i]
        fmt=self.w._number_formats[st.numFmtId-164] if st.numFmtId>=164 else st.numFmtId
        return (self.w._fonts[st.fontId],self.w._fills[st.fillId],self.w._borders[st.borderId],self.w._alignments[st.alignmentId],self.w._protections[st.protectionId],fmt)

a,b,p=map(Book,sys.argv[1:4]); issues=[]; style_issues=[]; errors_a=set();errors_b=set();dependent=[]
assert list(a.sheets)==list(b.sheets)
for row in range(3,135):
    assert b.cells['Carrying']['AO'+str(row)][1] == a.cells['Carrying']['AJ'+str(row)][1], ('Vendor',row)
for name,old in a.cells.items():
    new=b.cells[name]
    expected_positions=set()
    for addr,cell in old.items():
        dst=mapped(name,addr);expected_positions.add(dst)
        after=new.get(dst,(None,None,None,0));r,c=coordinate_to_tuple(dst)
        grid_formula=name=='Carrying' and 8<=r<=28 and c in (1,2,4,5,7,8,10,11,13,14,16,17,19,20,22,23,25,26,28,29,31,32)
        override=(name=='Docs' and dst=='E39') or (name=='Profit' and dst in ('L83','B43'))
        want=formula(cell[0],name) if cell[0] else cell[1]
        actual=after[0] if after[0] else after[1]
        if grid_formula: want=p.cells['Carrying'][dst][0].replace('VALUE(tblCarryingExpenses[Amount])','tblCarryingExpenses[Amount]')
        if name=='Docs' and dst=='E39':want='Profit!C9'
        if name=='Profit' and dst=='L83':want='IF(H83=0,0,+J83/H83*$F$75/DAYS(J75,H75))'
        if name=='Profit' and dst=='B43':want='+B28*SUM(B31:B42)'
        if name=='Carrying' and 3<=r<=134 and c==41: want=a.cells[name]['AJ'+str(r)][1]
        if name=='Carrying' and dst=='AQ86': want='38.78'
        if name=='Carrying' and dst=='AQ87': want='77.56'
        if name=='Carrying' and dst=='AV86': want=str(cell[1])+'; Grid override Carrying!Q12, 2026-09-30: 55.15 -> 38.78'
        if name=='Carrying' and dst=='AV87': want=str(cell[1])+'; Grid override Carrying!Q13, 2026-09-30: blank -> 77.56'
        addition=name=='Carrying' and c==34 and r<=29 and cell[0] is None and cell[1] is None
        same=formula_equal(want,actual) if cell[0] or grid_formula or override else want==actual
        if not same and not addition:issues.append([name,addr,dst,want,actual])
        if not(name=='Carrying' and c<=36 and r<=33) and a.style(cell[3])!=b.style(after[3]):style_issues.append([name,addr,dst])
        if cell[2]=='e':errors_a.add((name,dst,cell[1]))
        if cell[0] and cell[0]!=after[0] and name!='Carrying':dependent.append([name,addr,dst])
    for dst,cell in new.items():
        if cell[2]=='e':errors_b.add((name,dst,cell[1]))
        if dst not in expected_positions and (cell[0] is not None or cell[1] is not None):
            r,c=coordinate_to_tuple(dst)
            if not(name=='Carrying' and (r<=4 or c==50 or (c==41 and 3<=r<=134) or (34<=c<=36 and r<=29))) and not(name=='Profit' and r==41):issues.append(['new',name,dst])
    for merge in a.w[name].merged_cells.ranges:
        expected=':'.join(mapped(name,addr) for addr in str(merge).split(':'))
        if expected not in b.w[name].merged_cells:issues.append(['merge',name,str(merge),expected])
    for attr in ('page_setup','page_margins','print_options'):
        if getattr(a.w[name],attr)!=getattr(b.w[name],attr):issues.append([attr,name])
    for table_name,t in a.w[name].tables.items():
        t=a.w[name].tables[table_name]; u=b.w[name].tables.get(table_name)
        expected=':'.join(mapped(name,x) for x in t.ref.split(':'))
        if u is None or expected!=u.ref or [x.name for x in t.tableColumns]!=[x.name for x in u.tableColumns]:issues.append(['table',name,table_name])
for name,n in a.w.defined_names.items():
    after=b.w.defined_names.get(name)
    if after is None or not formula_equal(formula(n.attr_text,''),after.attr_text):issues.append(['name',name,n.attr_text,after.attr_text if after else None])
external=[n for n in b.z.namelist() if n.startswith('xl/externalLinks/')]
for k in ('A2','C2','G2','L2','R2','U2','A4','D4','I4','O4','U4','Z3'):
    if p.style(p.cells['Carrying'][k][3])!=b.style(b.cells['Carrying'][k][3]):issues.append(['form style',k])
def controls(book):
    return Counter(ET.tostring(ET.fromstring(book.z.read(n))) for n in book.z.namelist() if n.startswith('xl/ctrlProps/') and n.endswith('.xml'))
missing=list((controls(a)-controls(b)).elements())
result={'issues':issues,'style_issues':style_issues,'new_errors':sorted(errors_b-errors_a),'removed_errors':sorted(errors_a-errors_b),'existing_error_count':len(errors_a),'external_links':external,'missing_control_count':len(missing),'dependent_formula_change_count':len(dependent),'records':b.w['Carrying'].tables['tblCarryingExpenses'].ref}
print(json.dumps(result,indent=2))
if issues or style_issues or errors_b-errors_a or external or missing:sys.exit(1)
