"""Read-only current-source inventory for the full Carrying rollout."""
import collections
import datetime
import hashlib
import json
import pathlib
import sys
import warnings

import openpyxl
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

def val(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return to_excel(v)
    return v.text if hasattr(v, 'text') else v

def inspect(p):
    w = openpyxl.load_workbook(p)
    cache = openpyxl.load_workbook(p, data_only=True)
    out = dict(file=p.name, sha256=hashlib.sha256(p.read_bytes()).hexdigest(), sheets=w.sheetnames)
    out['profit'] = {c.coordinate: dict(formula=val(c.value), value=val(cache['Profit'][c.coordinate].value))
                     for row in w['Profit'] for c in row if c.value is not None}
    out['errors'] = {s.title: {c.coordinate: c.value for row in s for c in row if c.data_type == 'e'} for s in cache}
    if 'Carrying' not in w:
        out['missingSheet'] = True
        return out
    s, d = w['Carrying'], cache['Carrying']
    out['cells'] = {c.coordinate: dict(formula=val(c.value), value=val(d[c.coordinate].value), format=c.number_format)
                    for row in s for c in row if c.value is not None}
    out['names'] = {k: v.attr_text for k, v in w.defined_names.items() if k.startswith('ce')}
    out['merges'] = sorted(str(x) for x in s.merged_cells.ranges)
    if 'tblCarryingExpenses' not in s.tables:
        out['missingTable'] = True
        return out
    t = s.tables['tblCarryingExpenses']
    left, top, right, end = range_boundaries(t.ref)
    headers = [c.name for c in t.tableColumns]
    rows = [dict(zip(headers, [val(d.cell(r,c).value) for c in range(left,right+1)]), row=r,
                 formulas=[val(s.cell(r,c).value) for c in range(left,right+1)]) for r in range(top+1,end+1)]
    out.update(table=t.ref, rows=rows)
    hr = next(c.row for row in s.iter_rows(max_row=12, max_col=1) for c in row if c.value == 'Duke Electric')
    first = hr+3
    capacity = int(out['names']['ceDisplayCapacity'])
    footer = first+capacity
    out.update(header=hr, first=first, footer=footer, capacity=capacity)
    out['grid'] = []
    for col in range(1,40 if hr==7 else 37,3):
        cat = s.cell(hr,col).value
        internal = 'Mortgage Payment' if cat == 'Mortgage Payment paid after Reinstatement' else cat
        group = [r for r in rows if r['Category']==internal and r['Include']=='Yes']
        ordered = sorted(group, key=lambda r: (r['Date'] or 0, r['row']))
        overrides = []
        for r in range(first,footer):
            for c,field in ((col,'Date'),(col+1,'Amount')):
                cell = s.cell(r,c)
                if cell.value is not None and cell.data_type!='f' and not hasattr(cell.value,'text'):
                    overrides.append(dict(cell=cell.coordinate,field=field,value=val(cell.value),candidate=ordered[r-first] if r-first<len(ordered) else None))
        out['grid'].append(dict(category=internal,col=col,count=len(group),total=val(d.cell(footer,col+1).value),
                                totalFormula=val(s.cell(footer,col+1).value), overrides=overrides,
                                thirdColumn={s.cell(r,col+2).coordinate:val(s.cell(r,col+2).value) for r in range(hr,footer+1) if s.cell(r,col+2).value is not None}))
    out['zeros']=[r['row'] for r in rows if type(r['Amount']) in (int,float) and r['Amount']==0 and not str(r['formulas'][5]).startswith('=')]
    out['blankVendors']=[r['row'] for r in rows if not str(r['Vendor'] or '').strip()]
    return out

base = pathlib.Path(sys.argv[1])
out = {p.name:inspect(p) for p in sorted((base/'source').glob('*.xlsm'))}
(base/'inventory.json').write_text(json.dumps(out, indent=2, default=str), encoding='utf-8')
for name, x in out.items():
    print(json.dumps(dict(file=name, table=x.get('table'), rows=len(x.get('rows',[])), capacity=x.get('capacity'),
        zeros=x.get('zeros'),blankVendors=x.get('blankVendors'),grid=x.get('grid'),profitTotal=x['profit'].get('B43')),default=str))
