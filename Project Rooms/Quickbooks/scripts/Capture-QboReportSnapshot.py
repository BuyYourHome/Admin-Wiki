"""Capture immutable QBO report exports for the offline entity-reporting pilot.

This tool reads XLSX files without editing them, preserves exact source bytes in
an external SQLite archive, and audits input coverage. It does not connect to
QuickBooks, import transactions into books, or certify financial statements.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sqlite3
import sys


REPOSITORY = Path(__file__).resolve().parents[3]
SCHEMA_VERSION = 1
ARCHIVE_APPLICATION_ID = 0x51424F53  # QBOS; SQLite user_version alone is not an identity.
SNAPSHOT_COLUMNS = (
    "snapshot_id", "company", "source_path", "source_sha256", "source_bytes",
    "captured_utc", "content", "audit_json",
)
ALIASES = {
    "date": {"date", "transaction date"},
    "type": {"type", "transaction type"},
    "account": {"account", "account name", "distribution account"},
    "class": {"class", "class name"},
    "debit": {"debit", "debits"},
    "credit": {"credit", "credits"},
    "amount": {"amount"},
    "transaction_id": {"transaction id", "transactionid", "txn id", "txnid"},
}


def norm(value):
    return " ".join(str("" if value is None else value).strip().casefold().split())


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    # Bare numbers could be amounts or Excel serials without a date format.
    return None


def decimal_amount(value):
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, bool):
        raise ValueError("Boolean is not a posting amount")
    if isinstance(value, str):
        value = value.strip().replace(",", "").replace("$", "")
        if value.startswith("(") and value.endswith(")"):
            value = "-" + value[1:-1]
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("Unrecognized posting amount") from exc
    if not result.is_finite():
        raise ValueError("Non-finite posting amount")
    return result


def report_basis(metadata_rows):
    """Recognize standalone report metadata, never transaction memo text."""
    basis = set()
    for row in metadata_rows:
        values = [norm(value) for value in row if norm(value)]
        if len(values) != 1:
            continue
        match = re.fullmatch(
            r"(cash|accrual) basis(?: (?:monday|tuesday|wednesday|thursday|friday|saturday|sunday)\b.*)?",
            values[0],
        )
        if match:
            basis.add(match.group(1))
    return next(iter(basis)) if len(basis) == 1 else None


def inspect_rows(rows, uncached_formula_cells=()):
    """Audit a worksheet without assigning entities or inventing posting IDs."""
    result = {
        "row_count": len(rows), "column_count": max(map(len, rows), default=0),
        "header_row": None, "headers": [], "dated_row_count": 0,
        "months": {}, "transaction_types": {}, "classed_dated_rows": 0,
        "unclassed_dated_rows": 0, "transaction_ids_on_dated_rows": 0,
        "unaccounted_dated_rows": 0,
        "has_class_column": False, "has_debit_credit_columns": False,
        "has_explicit_account_column": False, "has_amount_column": False,
        "debit_credit_difference": None, "invalid_amount_rows": [],
        "missing_posting_amount_rows": [],
        "uncached_formula_cells": [
            {"row": row, "column": column}
            for row, column in sorted(set(uncached_formula_cells))
        ],
        "account_section_count": 0, "beginning_balance_row_count": 0,
        "accounting_basis": None, "coverage_findings": [],
        "entity_reporting_ready": False,
    }
    if result["uncached_formula_cells"]:
        result["coverage_findings"].append("uncached_formula_results")
    fields = {}
    for index, row in enumerate(rows[:50]):
        found = {}
        for col, cell in enumerate(row):
            for field, aliases in ALIASES.items():
                if norm(cell) in aliases:
                    found.setdefault(field, []).append(col)
        if "date" in found and "type" in found:
            if any(len(cols) > 1 for cols in found.values()):
                result["coverage_findings"].append("ambiguous_duplicate_header")
                return result
            fields = {field: cols[0] for field, cols in found.items()}
            result["header_row"] = index + 1
            result["headers"] = [str(cell or "") for cell in row]
            break
    if not fields:
        # Summary reports still carry recognized standalone basis metadata.
        # Finding it does not turn a summary into a transaction table.
        result["accounting_basis"] = report_basis(rows)
        result["coverage_findings"].append("no_transaction_table_detected")
        return result
    result["has_class_column"] = "class" in fields
    result["has_debit_credit_columns"] = {"debit", "credit"} <= fields.keys()
    result["has_explicit_account_column"] = "account" in fields
    result["has_amount_column"] = "amount" in fields
    dated = []
    last_dated_row = result["header_row"]
    months, types, sections = Counter(), Counter(), set()
    debit_total, credit_total = Decimal(0), Decimal(0)
    for row_number, row in enumerate(rows[result["header_row"]:], result["header_row"] + 1):
        def cell(field):
            col = fields.get(field)
            return row[col] if col is not None and col < len(row) else None

        values = [norm(value) for value in row if value is not None]
        if any(value in ("beginning balance", "opening balance") for value in values):
            result["beginning_balance_row_count"] += 1
        posting_date = parse_date(cell("date"))
        if posting_date is None:
            # Section headings are coverage hints, never authoritative account IDs.
            nonempty = [(i, value) for i, value in enumerate(row) if value not in (None, "")]
            if len(nonempty) == 1 and nonempty[0][0] < fields["date"]:
                label = str(nonempty[0][1]).strip()
                if label and not norm(label).startswith("total"):
                    sections.add(label)
            continue
        dated.append(posting_date)
        last_dated_row = row_number
        months[posting_date.strftime("%Y-%m")] += 1
        types[str(cell("type") or "[missing]")] += 1
        classed = bool(norm(cell("class")))
        result["classed_dated_rows"] += int(classed)
        result["unclassed_dated_rows"] += int(not classed)
        result["unaccounted_dated_rows"] += int(not bool(norm(cell("account"))))
        result["transaction_ids_on_dated_rows"] += int(bool(norm(cell("transaction_id"))))
        if result["has_debit_credit_columns"]:
            if not norm(cell("debit")) and not norm(cell("credit")):
                result["missing_posting_amount_rows"].append(row_number)
                continue
            try:
                debit_total += decimal_amount(cell("debit"))
                credit_total += decimal_amount(cell("credit"))
            except ValueError:
                result["invalid_amount_rows"].append(row_number)
    result["dated_row_count"] = len(dated)
    result["months"] = dict(sorted(months.items()))
    result["transaction_types"] = dict(sorted(types.items()))
    result["account_section_count"] = len(sections)
    result["date_min"] = min(dated).isoformat() if dated else None
    result["date_max"] = max(dated).isoformat() if dated else None
    # QBO can put a standalone basis/timestamp footer after the transaction body.
    metadata_rows = rows[:result["header_row"] - 1] + rows[last_dated_row:]
    result["accounting_basis"] = report_basis(metadata_rows)
    if (result["has_debit_credit_columns"] and dated
            and not result["invalid_amount_rows"]
            and not result["missing_posting_amount_rows"]
            and not result["uncached_formula_cells"]):
        result["debit_credit_difference"] = str(debit_total - credit_total)
    checks = {
        "no_dated_postings_detected": bool(dated),
        "missing_transaction_level_class": result["has_class_column"],
        "missing_explicit_debit_credit": result["has_debit_credit_columns"],
        "missing_explicit_account_column": result["has_explicit_account_column"],
        "missing_transaction_identifiers": bool(dated) and result["transaction_ids_on_dated_rows"] == len(dated),
        "accounting_basis_not_established": bool(result["accounting_basis"]),
    }
    result["coverage_findings"].extend(name for name, passed in checks.items() if not passed)
    if (result["debit_credit_difference"] is not None
            and Decimal(result["debit_credit_difference"]) != 0):
        result["coverage_findings"].append("dated_debits_and_credits_differ")
    if result["has_class_column"] and result["unclassed_dated_rows"]:
        result["coverage_findings"].append("missing_class_values")
    if result["has_explicit_account_column"] and result["unaccounted_dated_rows"]:
        result["coverage_findings"].append("missing_account_values")
    if result["missing_posting_amount_rows"]:
        result["coverage_findings"].append("missing_posting_amounts")
    if result["invalid_amount_rows"]:
        result["coverage_findings"].append("invalid_posting_amounts")
    result["coverage_findings"].append("entity_ownership_and_independent_controls_not_verified")
    # Structural presence can never certify legal ownership or book completeness.
    return result


def inspect_xlsx(content):
    import openpyxl

    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    formulas = None
    try:
        formulas = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheets = []
        for sheet, formula_sheet in zip(workbook.worksheets, formulas.worksheets):
            rows = list(sheet.iter_rows(values_only=True))
            uncached = []
            for row_number, row in enumerate(formula_sheet.iter_rows(), 1):
                for column, cell in enumerate(row, 1):
                    if cell.data_type == "f" and (
                        row_number > len(rows) or column > len(rows[row_number - 1])
                        or rows[row_number - 1][column - 1] is None
                    ):
                        uncached.append((row_number, column))
            sheets.append({"sheet": sheet.title, **inspect_rows(rows, uncached)})
        return {"sheets": sheets, "status": "raw_source_snapshot_only", "entity_reporting_ready": False}
    finally:
        if formulas is not None:
            formulas.close()
        workbook.close()


def validate_output(path, source_paths):
    if not path.is_absolute():
        raise ValueError("Output paths must be absolute")
    target = path.resolve()
    if target == REPOSITORY or REPOSITORY in target.parents:
        raise ValueError("Financial archives must remain outside the Git repository")
    if target in {source.resolve() for source in source_paths}:
        raise ValueError("An output cannot overwrite a source")
    return target


def open_archive(database):
    database.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(database))
    try:
        version = connection.execute("PRAGMA user_version").fetchone()[0]
        application_id = connection.execute("PRAGMA application_id").fetchone()[0]
        if version not in (0, SCHEMA_VERSION):
            raise ValueError("Unsupported archive schema version")
        tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if version == 0:
            if tables or application_id:
                raise ValueError("Refusing to initialize an existing unrelated database")
            connection.execute("""CREATE TABLE source_snapshots (
                snapshot_id TEXT PRIMARY KEY, company TEXT NOT NULL, source_path TEXT NOT NULL,
                source_sha256 TEXT NOT NULL, source_bytes INTEGER NOT NULL,
                captured_utc TEXT NOT NULL, content BLOB NOT NULL, audit_json TEXT NOT NULL,
                UNIQUE(company, source_sha256))""")
            connection.execute(f"PRAGMA application_id={ARCHIVE_APPLICATION_ID}")
            connection.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
            connection.commit()
        else:
            columns = tuple(row[1] for row in connection.execute("PRAGMA table_info(source_snapshots)"))
            if (application_id != ARCHIVE_APPLICATION_ID or tables != {"source_snapshots"}
                    or columns != SNAPSHOT_COLUMNS):
                raise ValueError("Unrecognized archive database identity or schema")
        return connection
    except Exception:
        connection.close()
        raise


def capture(connection, company, source):
    if not source.is_absolute() or source.suffix.casefold() != ".xlsx":
        raise ValueError("Sources must be absolute XLSX paths")
    before = source.stat()
    content = source.read_bytes()
    after = source.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise ValueError("Source changed while it was being read; capture again after export completes")
    digest = hashlib.sha256(content).hexdigest()
    identity = hashlib.sha256((company + "\0" + digest).encode("utf-8")).hexdigest()
    existing = connection.execute(
        "SELECT company, source_sha256, source_bytes, content, audit_json "
        "FROM source_snapshots WHERE snapshot_id=?", (identity,),
    ).fetchone()
    if existing:
        # Replay preserves the original cached inspection evidence. A newer
        # inspector must not silently rewrite this immutable snapshot's audit.
        archived_company, archived_digest, archived_size, archived_content, audit_json = existing
        if (archived_company != company or archived_digest != digest
                or not isinstance(archived_content, bytes)
                or archived_size != len(archived_content)
                or hashlib.sha256(archived_content).hexdigest() != digest):
            raise ValueError("Existing snapshot content integrity check failed")
        return {"snapshot_id": identity, "source_sha256": digest, "new_snapshot": False,
                "audit": json.loads(audit_json)}
    audit = inspect_xlsx(content)
    connection.execute("""INSERT INTO source_snapshots VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (
        identity, company, str(source.resolve()), digest, len(content),
        datetime.now(timezone.utc).isoformat(), content,
        json.dumps(audit, sort_keys=True, separators=(",", ":")),
    ))
    return {"snapshot_id": identity, "source_sha256": digest, "new_snapshot": True, "audit": audit}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--company", required=True)
    parser.add_argument("--database", type=Path, required=True)
    parser.add_argument("--source", type=Path, action="append", required=True)
    args = parser.parse_args(argv)
    if not args.company.strip():
        parser.error("Company must not be blank")
    database = validate_output(args.database, args.source)
    connection = open_archive(database)
    try:
        with connection:
            results = [capture(connection, args.company.strip(), source) for source in args.source]
        print(json.dumps({
            "database": str(database), "sources_processed": len(results),
            "new_snapshots": sum(item["new_snapshot"] for item in results),
            "source_audits": results, "entity_reporting_ready": False,
        }, indent=2))
    finally:
        connection.close()


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, sqlite3.Error) as exc:
        print(f"Capture failed: {exc}", file=sys.stderr)
        sys.exit(2)
